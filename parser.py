"""
Doküman ve slayt ayrıştırma modülü (PDF & PPTX).
PyMuPDF (fitz) ve python-pptx kullanarak sayfa/slayt bazlı anlamsal bloklar ve metadata üretir.
Gereksiz boşlukları ve gürültüleri temizleyerek token tasarrufu sağlar.
"""

import os
import re
from typing import List, Dict, Any

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None


def clean_text(text: str) -> str:
    """Metindeki gereksiz boşlukları, satır sonlarını ve fazlalıkları temizler."""
    if not text:
        return ""
    # Birden fazla ardışık boşluğu ve tab'ları teke indir
    text = re.sub(r"[ \t]+", " ", text)
    # 2'den fazla ardışık satır başlarını ikiye indir (paragraf yapısını koru)
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Satır başı ve sonundaki boşlukları kırp
    lines = [line.strip() for line in text.split("\n")]
    text = "\n".join(lines)
    return text.strip()


def parse_pdf(file_path: str) -> List[Dict[str, Any]]:
    """
    PDF dosyasını sayfa sayfa okur.
    Her sayfayı tek bir parça (chunk) olarak hazırlar.
    """
    if fitz is None:
        raise ImportError("PyMuPDF (fitz) kütüphanesi yüklü değil.")

    chunks = []
    file_name = os.path.basename(file_path)
    
    doc = fitz.open(file_path)
    for page_index in range(len(doc)):
        page = doc[page_index]
        text = page.get_text("text")
        cleaned = clean_text(text)
        
        # Eğer sayfa çok kısa (örn. sadece sayfa numarası veya boş) ise token harcamamak için atla
        if len(cleaned) < 30:
            continue

        chunks.append({
            "content": cleaned,
            "metadata": {
                "source": file_name,
                "file_type": "pdf",
                "page": page_index + 1,
                "total_pages": len(doc),
            }
        })
    doc.close()
    return chunks


def parse_pptx(file_path: str) -> List[Dict[str, Any]]:
    """
    PowerPoint (.pptx) sunumunu slayt slayt okur.
    Başlık ve metin kutularını sırayla birleştirir.
    """
    if Presentation is None:
        raise ImportError("python-pptx kütüphanesi yüklü değil.")

    chunks = []
    file_name = os.path.basename(file_path)
    
    prs = Presentation(file_path)
    total_slides = len(prs.slides)
    
    for slide_index, slide in enumerate(prs.slides):
        slide_texts = []
        
        # Slayttaki tüm şekilleri (şekil, metin kutusu, tablo) tara
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
            elif shape.has_table:
                # Tablo hücrelerini oku
                for row in shape.table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        slide_texts.append(" | ".join(row_texts))

        raw_text = "\n".join(slide_texts)
        cleaned = clean_text(raw_text)

        if len(cleaned) < 25:
            continue

        chunks.append({
            "content": cleaned,
            "metadata": {
                "source": file_name,
                "file_type": "pptx",
                "page": slide_index + 1,  # Slayt numarası
                "total_pages": total_slides,
            }
        })
        
    return chunks


def parse_document(file_path: str) -> List[Dict[str, Any]]:
    """Dosya uzantısına göre uygun ayrıştırıcıyı çağırır."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".pptx":
        return parse_pptx(file_path)
    else:
        print(f"[UYARI] Desteklenmeyen dosya türü: {file_path}")
        return []
